import re
import fitz
from docx import Document
from io import BytesIO


# Technical skills our parser can recognize
KNOWN_SKILLS = [
    "Java", "Python", "JavaScript", "TypeScript",
    "Spring Boot", "Spring MVC", "Spring",
    "React.js", "React", "Node.js",
    "HTML", "CSS", "Tailwind CSS",
    "MySQL", "PostgreSQL", "MongoDB", "SQL",
    "Hibernate", "JPA", "JDBC",
    "REST APIs", "REST API", "JWT",
    "Git", "GitHub", "Docker",
    "AWS", "Azure", "Linux",
    "Data Structures and Algorithms", "DSA",
    "OOP", "DBMS", "Maven", "Microservices",
    "C", "C++", "C#", "Java Collections",
    "Multithreading", "Stream API"
]


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF resume."""
    text = ""

    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            text += page.get_text() + "\n"

    return text.strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX resume."""
    document = Document(BytesIO(file_bytes))
    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()


def extract_contact_details(text: str) -> dict:
    """Extract email address and Indian mobile number."""
    email_match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    )

    phone_match = re.search(
        r"(?<!\d)(?:\+91[\s-]?)?[6-9]\d{9}(?!\d)",
        text
    )

    return {
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0) if phone_match else None
    }


def extract_name(text: str) -> str | None:
    """Guess the candidate's name from the beginning of the resume."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    for line in lines[:8]:
        if (
            len(line) <= 60
            and re.fullmatch(r"[A-Za-z]+(?:[ .'-][A-Za-z]+)*", line)
            and not re.search(
                r"\b(resume|curriculum vitae|developer|engineer|"
                r"student|contact|profile|summary)\b",
                line,
                re.IGNORECASE
            )
        ):
            return line

    return None


def extract_designation(text: str) -> str | None:
    """Identify a likely professional designation."""
    designations = [
        "Full Stack Java Developer",
        "Full Stack Developer",
        "Java Developer",
        "Python Developer",
        "Software Engineer",
        "Software Developer",
        "Frontend Developer",
        "Backend Developer",
        "Web Developer",
        "Software Developer Intern",
        "Java Developer Intern"
    ]

    for designation in designations:
        if re.search(re.escape(designation), text, re.IGNORECASE):
            return designation

    return None


def extract_skills(text: str) -> list:
    """Find known technical skills mentioned in the resume."""
    found_skills = []
    text_lower = text.lower()

    for skill in KNOWN_SKILLS:
        pattern = r"(?<![a-z0-9+#.])" + re.escape(skill.lower()) + r"(?![a-z0-9+#.])"

        if re.search(pattern, text_lower):
            # Avoid returning the same skill twice.
            normalized = skill.lower()

            if normalized not in [item.lower() for item in found_skills]:
                found_skills.append(skill)

    return found_skills


def extract_education(text: str) -> list:
    """Extract lines that appear to describe education."""
    education = []
    lines = text.splitlines()

    keywords = [
        "b.tech", "bachelor", "b.e.", "b.e ",
        "m.tech", "master", "b.sc", "m.sc",
        "bca", "mca", "diploma", "12th", "10th",
        "computer science", "university", "college"
    ]

    for line in lines:
        cleaned = line.strip()

        if cleaned and any(
            keyword in cleaned.lower() for keyword in keywords
        ):
            if cleaned not in education:
                education.append(cleaned)

    return education[:10]


def extract_projects(text: str) -> list:
    """Extract likely project names from project-section lines."""
    projects = []
    in_project_section = False

    section_pattern = re.compile(
        r"^\s*(projects|academic projects|personal projects|"
        r"key projects|project experience)\s*:?\s*$",
        re.IGNORECASE
    )

    stop_pattern = re.compile(
        r"^\s*(education|technical skills|skills|certifications|"
        r"experience|internship|achievements|languages)\s*:?\s*$",
        re.IGNORECASE
    )

    for line in text.splitlines():
        cleaned = line.strip()

        if section_pattern.match(cleaned):
            in_project_section = True
            continue

        if stop_pattern.match(cleaned):
            in_project_section = False
            continue

        if in_project_section and cleaned:
            # Ignore bullet-only lines and very long descriptions.
            cleaned = re.sub(r"^[•\-*▪]\s*", "", cleaned)

            if (
                cleaned
                and len(cleaned) <= 100
                and cleaned not in projects
            ):
                projects.append(cleaned)

    return projects[:10]


def extract_experience(text: str) -> str | None:
    """Find a likely experience duration or identify a fresher."""
    patterns = [
        r"\b\d+(?:\.\d+)?\+?\s+years?\s+(?:of\s+)?experience\b",
        r"\b\d+(?:\.\d+)?\+?\s+months?\s+(?:of\s+)?experience\b"
    ]

    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)

        if match:
            return match.group(0)

    if re.search(r"\bfresher\b|\brecent graduate\b", text, re.IGNORECASE):
        return "Fresher"

    if re.search(r"\bintern\b|\binternship\b", text, re.IGNORECASE):
        return "Internship experience mentioned; duration not identified"

    return None


def parse_resume(file_bytes: bytes, filename: str) -> dict:
    """Extract structured information from a PDF or DOCX resume."""
    extension = (
        filename.lower().rsplit(".", 1)[-1]
        if "." in filename else ""
    )

    if extension == "pdf":
        text = extract_text_from_pdf(file_bytes)

    elif extension == "docx":
        text = extract_text_from_docx(file_bytes)

    else:
        raise ValueError(
            "Unsupported file format. Please upload a PDF or DOCX resume."
        )

    if not text.strip():
        raise ValueError(
            "No readable text found. Please upload a text-based PDF or DOCX resume."
        )

    contact_details = extract_contact_details(text)

    return {
        "filename": filename,
        "name": extract_name(text),
        "email": contact_details["email"],
        "phone": contact_details["phone"],
        "designation": extract_designation(text),
        "skills": extract_skills(text),
        "experience": extract_experience(text),
        "education": extract_education(text),
        "projects": extract_projects(text),
        "raw_text": text
    }